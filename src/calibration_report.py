"""Lightbox calibration report: measured light field across the test plane.

Takes one calibration campaign folder (calibration/cal-...) and builds, from
its captures.csv, the five outputs on AK's 2026-09-25 sheet:

  1. VAC vs max lux and max irradiance (table + data for graphs)
  2. Normalised heatmaps per level: % of that level's max, 0-100 % -> 0-255 gray
  3. (dropped 2026-09-28: no curve fitting, the fit errors were too large)
  4. Lux vs irradiance pairs for every node at every level
  5. Panel averages of irr and lux over a w x l mm panel centred on F4, at the
     measured VAC levels, and the reverse: the VAC for a required panel-average
     irradiance (power-law interpolation between the two bracketing levels)

Everything is recomputed from the selected folder, so a new campaign needs no
code change.

Decisions (see the campaign notes.md):
  * Irradiance zero offset of 0.23 W/m2 is subtracted from every irr reading.
  * Values are NOT drift-corrected. Each pass's reference drift
    (ref_end - ref_start) / ref_start is reported as an uncertainty.
  * No curve fitting. Panel averages come straight from the measured map:
    the map is interpolated bilinearly between nodes and averaged over the
    panel rectangle on a 5 mm grid. Only measured VAC levels are available.
  * Single-node glitches (value more than 50 % away from the median of its
    neighbours, or <= 0 after the offset) are flagged and left out of the max.
    For panel averages only, a glitch node is replaced by that neighbour median.

Usage:
  python src/calibration_report.py calibration/cal-20260917-113349-lightbox
  python src/calibration_report.py <folder> --vac 54 --panel 200x150
  python src/calibration_report.py <folder> --target-irr 100 --panel 200x150
"""
from __future__ import annotations

import argparse
import json
import math
import sys
import time
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
import calibration_model as cm  # noqa: E402

IRR_ZERO_OFFSET_WM2 = 0.23
OUTLIER_FRAC = 0.5
PANEL_INTEGRATION_STEP_MM = 5.0
CONTOUR_LEVELS_PCT = [50, 60, 70, 80, 90, 95]
QUANTITIES = ("lux", "irr")
UNITS = {"lux": "lx", "irr": "W/m2"}
STANDARD_PANELS_MM = [(100, 100), (150, 150), (200, 150), (250, 200), (300, 200), (400, 300), (500, 300)]
REPORT_SUBDIR = "report"


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------


def load_passes(folder, irr_offset: float = IRR_ZERO_OFFSET_WM2):
    """Return (config, grid, passes). passes[(level_index, quantity)] holds the
    raw node matrix (rows x cols, offset applied for irr), 3-sigma %, refs and
    drift. Passes missing a start or end reference are skipped."""
    folder = Path(folder)
    config = cm.load_campaign_config(folder)
    grid = cm.grid_from_campaign(config)
    caps = cm.active_captures(cm.read_captures(cm.captures_csv_path(folder)))

    groups: dict[tuple[int, str], list[dict]] = {}
    for c in caps:
        groups.setdefault((c["level_index"], c["quantity"]), []).append(c)

    passes = {}
    for (level, q), rows in sorted(groups.items()):
        offset = irr_offset if q == "irr" else 0.0
        ref_s = [r for r in rows if r["capture_kind"] == "ref_start" and r["mean"] is not None]
        ref_e = [r for r in rows if r["capture_kind"] == "ref_end" and r["mean"] is not None]
        if not ref_s or not ref_e:
            continue
        vals = np.full((grid.rows, grid.cols), np.nan)
        tsig = np.full((grid.rows, grid.cols), np.nan)
        for r in rows:
            if r["capture_kind"] != "grid" or r["mean"] is None:
                continue
            vals[r["row_index"], r["col_index"]] = r["mean"] - offset
            tsig[r["row_index"], r["col_index"]] = r["three_sigma_pct"] if r["three_sigma_pct"] is not None else np.nan
        rs = ref_s[-1]["mean"] - offset
        re_ = ref_e[-1]["mean"] - offset
        passes[(level, q)] = {
            "level": level,
            "quantity": q,
            "vac": float(rows[0]["variac_vac"]),
            "values": vals,
            "three_sigma": tsig,
            "ref_start": rs,
            "ref_end": re_,
            "drift_pct": (re_ - rs) / rs * 100.0 if rs else float("nan"),
            "missing_nodes": int(np.isnan(vals).sum()),
        }
    return config, grid, passes


def flag_outliers(vals: np.ndarray, frac: float = OUTLIER_FRAC):
    """Mask of single-node glitches, plus the neighbour median used."""
    R, C = vals.shape
    mask = np.zeros_like(vals, dtype=bool)
    med = np.full_like(vals, np.nan)
    for r in range(R):
        for c in range(C):
            v = vals[r, c]
            if not np.isfinite(v):
                continue
            nb = [
                vals[rr, cc]
                for rr in range(max(0, r - 1), min(R, r + 2))
                for cc in range(max(0, c - 1), min(C, c + 2))
                if (rr, cc) != (r, c) and np.isfinite(vals[rr, cc])
            ]
            if len(nb) < 2:
                continue
            m = float(np.median(nb))
            med[r, c] = m
            if v <= 0 or (m > 0 and abs(v - m) / m > frac):
                mask[r, c] = True
    return mask, med



# ---------------------------------------------------------------------------
# Panel averages from the measured map (no fitting)
# ---------------------------------------------------------------------------


def grid_centre(grid: cm.GridConfig):
    return (grid.cols - 1) * grid.pitch_mm / 2.0, (grid.rows - 1) * grid.pitch_mm / 2.0


def covered_mask(grid, w_mm, l_mm):
    m = np.zeros((grid.rows, grid.cols), dtype=bool)
    for lab in cm.covered_nodes(grid, w_mm, l_mm):
        c, r = cm.parse_node_label(lab)
        m[r, c] = True
    return m


def filled_map(vals: np.ndarray, mask: np.ndarray) -> np.ndarray:
    """Measured map with glitch or missing nodes replaced by the median of
    their good neighbours. Used for area averaging only."""
    out = vals.copy()
    bad = mask | ~np.isfinite(vals)
    R, C = vals.shape
    for r, c in np.argwhere(bad):
        nb = [vals[rr, cc] for rr in range(max(0, r - 1), min(R, r + 2)) for cc in range(max(0, c - 1), min(C, c + 2))
              if (rr, cc) != (r, c) and not bad[rr, cc]]
        out[r, c] = float(np.median(nb)) if nb else np.nan
    return out


def bilinear(m: np.ndarray, pitch_mm: float, x_mm, y_mm):
    """Bilinear interpolation of a node map (rows x cols) at x, y in mm from A1."""
    R, C = m.shape
    fx = np.clip(np.asarray(x_mm, float) / pitch_mm, 0, C - 1)
    fy = np.clip(np.asarray(y_mm, float) / pitch_mm, 0, R - 1)
    c0 = np.minimum(np.floor(fx).astype(int), C - 2)
    r0 = np.minimum(np.floor(fy).astype(int), R - 2)
    tx, ty = fx - c0, fy - r0
    return (m[r0, c0] * (1 - tx) * (1 - ty) + m[r0, c0 + 1] * tx * (1 - ty)
            + m[r0 + 1, c0] * (1 - tx) * ty + m[r0 + 1, c0 + 1] * tx * ty)


def check_panel(grid, w_mm, l_mm):
    x_c, y_c = grid_centre(grid)
    if w_mm <= 0 or l_mm <= 0:
        raise ValueError("panel width and length must be > 0 mm")
    if w_mm > 2 * x_c + 1e-9 or l_mm > 2 * y_c + 1e-9:
        raise ValueError(f"panel {w_mm:g} x {l_mm:g} mm is larger than the measured grid ({2*x_c:g} x {2*y_c:g} mm).")


def panel_average(filled: np.ndarray, grid, w_mm, l_mm, step_mm=PANEL_INTEGRATION_STEP_MM) -> float:
    """Area average of the measured map over a w x l mm rectangle centred on
    F4. w runs along the columns (A to K), l along the rows (1 to 7)."""
    check_panel(grid, w_mm, l_mm)
    x_c, y_c = grid_centre(grid)
    nx, ny = max(1, math.ceil(w_mm / step_mm)), max(1, math.ceil(l_mm / step_mm))
    xs = x_c - w_mm / 2 + (np.arange(nx) + 0.5) * w_mm / nx
    ys = y_c - l_mm / 2 + (np.arange(ny) + 0.5) * l_mm / ny
    X, Y = np.meshgrid(xs, ys)
    return float(np.mean(bilinear(filled, grid.pitch_mm, X, Y)))


def _loglog(x0, y0, x1, y1, y):
    """x at which the power law through (x0, y0) and (x1, y1) reaches y."""
    return math.exp(math.log(x0) + (math.log(y) - math.log(y0)) * (math.log(x1) - math.log(x0)) / (math.log(y1) - math.log(y0)))


def vac_for_target(passes, grid, w_mm, l_mm, target, quantity="irr"):
    """VAC that gives a panel-average `target` over a w x l panel centred on F4.

    Uses only measured data: the panel average at every measured level (same
    bilinear area average as panel_average), then a power-law (log-log)
    interpolation between the two levels that bracket the target. Raises
    ValueError outside the measured range. Returns a dict with the VAC, its
    uncertainty from drift, the bracketing levels, the local exponent n and the
    expected panel average of the other quantity at that VAC."""
    check_panel(grid, w_mm, l_mm)
    if target <= 0:
        raise ValueError("target must be > 0")
    other = "lux" if quantity == "irr" else "irr"
    levels = sorted({l for (l, q) in passes})
    pts = sorted(
        (passes[(l, quantity)]["vac"], panel_average(passes[(l, quantity)]["filled"], grid, w_mm, l_mm),
         abs(passes[(l, quantity)]["drift_pct"]), panel_average(passes[(l, other)]["filled"], grid, w_mm, l_mm), l)
        for l in levels
    )
    avgs = [p[1] for p in pts]
    if any(b <= a for a, b in zip(avgs, avgs[1:])):
        raise ValueError("panel average does not rise steadily with VAC for this panel size; check the data before inverting")
    if not (avgs[0] <= target <= avgs[-1]):
        raise ValueError(f"target {target:g} {UNITS[quantity]} is outside the measured range for a {w_mm:g} x {l_mm:g} mm panel "
                         f"({avgs[0]:.4g} to {avgs[-1]:.4g} {UNITS[quantity]}, {pts[0][0]:g} to {pts[-1][0]:g} VAC).")
    i = next(k for k in range(len(pts) - 1) if pts[k][1] <= target <= pts[k + 1][1])
    lo, hi = pts[i], pts[i + 1]
    vac = _loglog(lo[0], lo[1], hi[0], hi[1], target)
    n = math.log(hi[1] / lo[1]) / math.log(hi[0] / lo[0])
    drift = max(lo[2], hi[2])
    other_avg = math.exp(math.log(lo[3]) + (math.log(vac) - math.log(lo[0])) * (math.log(hi[3]) - math.log(lo[3])) / (math.log(hi[0]) - math.log(lo[0])))
    return {
        "vac": vac, "vac_unc_pct": drift / n, "vac_unc_v": vac * drift / n / 100.0, "drift_pct": drift, "n": n,
        "bracket": [{"vac": lo[0], "avg": lo[1], "level": lo[4]}, {"vac": hi[0], "avg": hi[1], "level": hi[4]}],
        f"{other}_avg": other_avg,
    }


# ---------------------------------------------------------------------------
# Report assembly
# ---------------------------------------------------------------------------


def contour_segments(pct: np.ndarray, levels=CONTOUR_LEVELS_PCT):
    """Contour lines of a % map, in node-index coordinates (col, row)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots()
    try:
        cs = ax.contour(np.arange(pct.shape[1]), np.arange(pct.shape[0]), pct, levels=levels)
        out = []
        for lvl, segs in zip(cs.levels, cs.allsegs):
            for seg in segs:
                if len(seg) >= 2:
                    out.append({"level": float(lvl), "pts": [[round(float(a), 3), round(float(b), 3)] for a, b in seg]})
        return out
    finally:
        plt.close(fig)


def _r(x, nd=4):
    if x is None:
        return None
    try:
        if not np.isfinite(x):
            return None
    except TypeError:
        return None
    return float(round(float(x), nd)) if abs(x) >= 1e-3 or x == 0 else float(f"{x:.4g}")


def _mat(m, nd=4):
    return [[_r(v, nd) for v in row] for row in m]


def build_report(folder, irr_offset=IRR_ZERO_OFFSET_WM2, outlier_frac=OUTLIER_FRAC):
    folder = Path(folder)
    config, grid, passes = load_passes(folder, irr_offset)
    for p in passes.values():
        p["outliers"], p["neighbour_median"] = flag_outliers(p["values"], outlier_frac)
        p["filled"] = filled_map(p["values"], p["outliers"])

    levels = sorted({l for (l, q) in passes if (l, "lux") in passes and (l, "irr") in passes})
    if not levels:
        raise SystemExit("No level in this campaign has both a complete lux and irr pass.")
    passes = {k: v for k, v in passes.items() if k[0] in levels}
    vac_by_level = {l: passes[(l, "lux")]["vac"] for l in levels}
    labels = [[cm.node_label(c, r) for c in range(grid.cols)] for r in range(grid.rows)]
    x_c, y_c = grid_centre(grid)
    cr, cc = int(round(y_c / grid.pitch_mm)), int(round(x_c / grid.pitch_mm))

    level_out = []
    for l in levels:
        entry = {"level": l, "vac": vac_by_level[l]}
        for q in QUANTITIES:
            p = passes[(l, q)]
            vals, mask = p["values"], p["outliers"]
            good = np.isfinite(vals) & ~mask
            vmax = float(np.max(vals[good]))
            rmax, cmax = np.argwhere((vals == vmax) & good)[0]
            pct = vals / vmax * 100.0
            gray = np.clip(np.rint(pct / 100.0 * 255.0), 0, 255)
            entry[q] = {
                "max": _r(vmax), "max_node": labels[rmax][cmax],
                "centre_value": _r(vals[cr, cc]),
                "drift_pct": _r(p["drift_pct"], 3),
                "ref_start": _r(p["ref_start"]), "ref_end": _r(p["ref_end"]),
                "median_three_sigma_pct": _r(float(np.nanmedian(p["three_sigma"])), 3),
                "missing_nodes": p["missing_nodes"],
                "values": _mat(vals), "filled": _mat(p["filled"]), "pct": _mat(pct, 2),
                "gray": [[None if not np.isfinite(g) else int(g) for g in row] for row in gray],
                "outliers": [labels[r][c] for r, c in np.argwhere(mask)],
                "contours": contour_segments(np.where(mask | ~np.isfinite(pct), np.nan, pct)),
            }
        level_out.append(entry)

    scatter = []
    for l in levels:
        pl, pi = passes[(l, "lux")], passes[(l, "irr")]
        for r in range(grid.rows):
            for c in range(grid.cols):
                lx, ir = pl["values"][r, c], pi["values"][r, c]
                if not (np.isfinite(lx) and np.isfinite(ir)):
                    continue
                scatter.append({
                    "node": labels[r][c], "level": l, "vac": vac_by_level[l], "lux": _r(lx), "irr": _r(ir),
                    "dist_mm": _r(math.hypot(c * grid.pitch_mm - x_c, r * grid.pitch_mm - y_c), 1),
                    "outlier": bool(pl["outliers"][r, c] or pi["outliers"][r, c]),
                })

    report = {
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "campaign": {
            "campaign_id": config.get("campaign_id"), "operator": config.get("operator"),
            "created_at": config.get("created_at"), "folder": folder.name,
            "grid": {"cols": grid.cols, "rows": grid.rows, "pitch_mm": grid.pitch_mm},
            "centre_node": labels[cr][cc], "x_c": x_c, "y_c": y_c,
            "dwell_seconds": config.get("dwell_seconds"), "warmup_seconds": config.get("warmup_seconds"),
            "sensor_head_height_mm": config.get("sensor_head_height_mm"),
        },
        "settings": {
            "irr_zero_offset_wm2": irr_offset, "drift_corrected": False, "curve_fit": False,
            "panel_average": f"bilinear between nodes, averaged on a {PANEL_INTEGRATION_STEP_MM:g} mm grid",
            "outlier_rule": f"|value - median(neighbours)| / median > {outlier_frac:g}, or value <= 0 after offset",
            "contour_levels_pct": CONTOUR_LEVELS_PCT,
        },
        "col_labels": [cm.col_letter(i) for i in range(grid.cols)],
        "levels": level_out,
        "scatter": scatter,
        "standard_panels_mm": STANDARD_PANELS_MM,
    }
    report["panel_table"] = [panel_row(grid, passes, l, vac_by_level[l], w, ln) for (w, ln) in STANDARD_PANELS_MM for l in levels]
    return report, passes, grid


def panel_row(grid, passes, level, vac, w, l):
    row = {"w_mm": w, "l_mm": l, "vac": vac, "level": level}
    for q in QUANTITIES:
        p = passes[(level, q)]
        m = covered_mask(grid, w, l) & np.isfinite(p["values"]) & ~p["outliers"]
        row[f"{q}_area_avg"] = _r(panel_average(p["filled"], grid, w, l))
        row[f"{q}_nodes_mean"] = _r(p["values"][m].mean()) if m.any() else None
        row[f"{q}_n_nodes"] = int(m.sum())
        row[f"{q}_drift_pct"] = _r(abs(p["drift_pct"]), 2)
    return row


def level_for_vac(passes, vac, tol=0.25):
    by_vac = {p["vac"]: p["level"] for (l, q), p in passes.items()}
    hit = [lv for v, lv in by_vac.items() if abs(v - vac) <= tol]
    if not hit:
        raise ValueError(f"{vac:g} VAC was not measured. Measured levels: {', '.join(f'{v:g}' for v in sorted(by_vac))} VAC. "
                         "There is no curve fit, so only measured levels can be used.")
    return hit[0]


# ---------------------------------------------------------------------------
# Outputs
# ---------------------------------------------------------------------------

def write_heatmap_pngs(report, out_dir: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    cols = report["col_labels"]
    rows = report["campaign"]["grid"]["rows"]
    for e in report["levels"]:
        for q in QUANTITIES:
            d = e[q]
            gray = np.array([[np.nan if g is None else g for g in r] for r in d["gray"]], dtype=float)
            pct = np.array([[np.nan if g is None else g for g in r] for r in d["pct"]], dtype=float)
            fig, ax = plt.subplots(figsize=(9, 5.2), dpi=130)
            ax.imshow(gray, cmap="gray", vmin=0, vmax=255, interpolation="nearest")
            masked = np.ma.masked_invalid(np.where([[ (cols[c] + str(r + 1)) in d["outliers"] for c in range(len(cols))] for r in range(rows)], np.nan, pct))
            cs = ax.contour(np.arange(len(cols)), np.arange(rows), masked, levels=CONTOUR_LEVELS_PCT, colors="#d95926", linewidths=1.1)
            ax.clabel(cs, fmt="%g%%", fontsize=7, inline=True, inline_spacing=2)
            for r in range(rows):
                for c in range(len(cols)):
                    if np.isfinite(pct[r, c]):
                        lab = cols[c] + str(r + 1)
                        txt = f"{pct[r, c]:.0f}" + (" x" if lab in d["outliers"] else "")
                        ax.text(c + 0.32, r + 0.34, txt, ha="right", va="bottom", fontsize=6.5,
                                color="#e34948" if lab in d["outliers"] else ("black" if gray[r, c] > 128 else "white"))
            ax.set_xticks(range(len(cols)), cols)
            ax.set_yticks(range(rows), [str(i + 1) for i in range(rows)])
            ax.set_title(f"L{e['level']} {e['vac']:g} VAC  {q}  % of max ({d['max']:.4g} {UNITS[q]} at {d['max_node']}), "
                         f"drift {d['drift_pct']:+.1f}%  (0-100% = gray 0-255)", fontsize=9)
            fig.tight_layout()
            fig.savefig(out_dir / f"heatmap_L{e['level']:02d}_{int(e['vac'])}VAC_{q}.png")
            plt.close(fig)



def write_workbook(report, path: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("openpyxl not installed; skipping the Excel workbook (pip install openpyxl)")
        return False
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "README"
    s = report["settings"]
    for line in [
        f"Campaign: {report['campaign']['campaign_id']}", f"Generated: {report['generated_at']}",
        f"Irradiance zero offset subtracted: {s['irr_zero_offset_wm2']} W/m2", "Drift correction: none (drift reported as uncertainty)",
        "Curve fitting: none. Panel averages use the measured map only.",
        f"Panel average: {s['panel_average']}. Glitch nodes replaced by their neighbour median for this only.",
        f"Outlier rule: {s['outlier_rule']}",
        "x runs A to K, y runs row 1 to row 7, origin A1, F4 = centre. Panel w is along x, l along y.",
    ]:
        ws.append([line])

    ws = wb.create_sheet("1_Max_vs_VAC")
    ws.append(["level", "VAC", "max lux", "lux node", "lux drift %", "max irr W/m2", "irr node", "irr drift %", "F4 lux", "F4 irr"])
    for e in report["levels"]:
        ws.append([e["level"], e["vac"], e["lux"]["max"], e["lux"]["max_node"], e["lux"]["drift_pct"],
                   e["irr"]["max"], e["irr"]["max_node"], e["irr"]["drift_pct"], e["lux"]["centre_value"], e["irr"]["centre_value"]])

    cols = report["col_labels"]
    for sheet, key in (("2_Heatmap_pct", "pct"), ("2_Heatmap_gray255", "gray"), ("Raw_values", "values")):
        ws = wb.create_sheet(sheet)
        for e in report["levels"]:
            for q in QUANTITIES:
                ws.append([f"L{e['level']} {e['vac']:g} VAC {q}", f"max {e[q]['max']} {UNITS[q]} at {e[q]['max_node']}", f"outliers: {', '.join(e[q]['outliers']) or 'none'}"])
                ws.cell(ws.max_row, 1).font = bold
                ws.append([""] + cols)
                for r, row in enumerate(e[q][key]):
                    ws.append([r + 1] + row)
                    if key == "gray":
                        for c, g in enumerate(row):
                            if g is not None:
                                ws.cell(ws.max_row, c + 2).fill = PatternFill("solid", fgColor=f"{g:02X}" * 3)
                                ws.cell(ws.max_row, c + 2).font = Font(color="000000" if g > 128 else "FFFFFF")
                ws.append([])

    ws = wb.create_sheet("4_Lux_vs_Irr")
    ws.append(["node", "level", "VAC", "lux", "irr W/m2", "distance from F4 mm", "outlier"])
    for p in report["scatter"]:
        ws.append([p["node"], p["level"], p["vac"], p["lux"], p["irr"], p["dist_mm"], int(p["outlier"])])

    ws = wb.create_sheet("5_Panel_avg")
    ws.append(["w mm", "l mm", "VAC", "avg irr (area)", "irr nodes mean", "irr n nodes", "irr +/- drift %",
               "avg lux (area)", "lux nodes mean", "lux n nodes", "lux +/- drift %"])
    for r in report["panel_table"]:
        ws.append([r["w_mm"], r["l_mm"], r["vac"], r["irr_area_avg"], r["irr_nodes_mean"], r["irr_n_nodes"], r["irr_drift_pct"],
                   r["lux_area_avg"], r["lux_nodes_mean"], r["lux_n_nodes"], r["lux_drift_pct"]])

    ws = wb.create_sheet("Drift_Outliers")
    ws.append(["level", "VAC", "quantity", "ref start", "ref end", "drift %", "median 3sigma %", "outlier nodes"])
    for e in report["levels"]:
        for q in QUANTITIES:
            d = e[q]
            ws.append([e["level"], e["vac"], q, d["ref_start"], d["ref_end"], d["drift_pct"], d["median_three_sigma_pct"], ", ".join(d["outliers"])])

    for sh in wb.worksheets:
        for cell in sh[1]:
            cell.font = bold
    wb.save(path)
    return True


def write_html(report, path: Path):
    """Standalone review page: the template next to this script with the
    report data inlined. Opens in any browser, no server needed."""
    template = Path(__file__).resolve().parent / "calibration_report_template.html"
    if not template.exists():
        print(f"{template.name} not found; skipping report.html")
        return False
    data = json.dumps(report, separators=(",", ":")).replace("</", "<\\/")
    body = template.read_text(encoding="utf-8").replace("__DATA__", data)
    path.write_text('<!doctype html><html lang="en"><head><meta charset="utf-8">'
                    '<meta name="viewport" content="width=device-width, initial-scale=1"></head><body>'
                    + body + "</body></html>", encoding="utf-8")
    return True



def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", help="calibration campaign folder (contains campaign.json and captures.csv)")
    ap.add_argument("--vac", type=float, help="query: a measured variac level, VAC")
    ap.add_argument("--panel", help="query: panel size WxL in mm, e.g. 200x150 (W along A-K, L along rows)")
    ap.add_argument("--target-irr", type=float, help="query: required panel-average irradiance, W/m2 (needs --panel); prints the VAC")
    ap.add_argument("--irr-offset", type=float, default=IRR_ZERO_OFFSET_WM2)
    ap.add_argument("--no-png", action="store_true")
    ap.add_argument("--no-xlsx", action="store_true")
    args = ap.parse_args(argv)

    report, passes, grid = build_report(args.folder, args.irr_offset)

    if args.target_irr is not None:
        if not args.panel:
            ap.error("--target-irr needs --panel")
        w, l = (float(v) for v in args.panel.lower().split("x"))
        try:
            r = vac_for_target(passes, grid, w, l, args.target_irr, "irr")
        except ValueError as exc:
            print(exc)
            return 2
        b0, b1 = r["bracket"]
        print(f"{args.target_irr:g} W/m2 average over {w:g} x {l:g} mm -> {r['vac']:.1f} VAC  +/- {r['vac_unc_v']:.2f} V "
              f"(drift {r['drift_pct']:.1f}% / n {r['n']:.2f})")
        print(f"  between measured {b0['vac']:g} VAC ({b0['avg']:.4g} W/m2) and {b1['vac']:g} VAC ({b1['avg']:.4g} W/m2); "
              f"expected avg lux {r['lux_avg']:.4g} lx")
        return 0

    if args.vac is not None or args.panel:
        if args.vac is None or not args.panel:
            ap.error("--vac and --panel go together")
        w, l = (float(v) for v in args.panel.lower().split("x"))
        try:
            level = level_for_vac(passes, args.vac)
            for q in ("irr", "lux"):
                p = passes[(level, q)]
                avg = panel_average(p["filled"], grid, w, l)
                print(f"avg {q} over {w:g} x {l:g} mm at {p['vac']:g} VAC (L{level}) = {avg:.4g} {UNITS[q]}  +/- {abs(p['drift_pct']):.1f}% (drift)")
        except ValueError as exc:
            print(exc)
            return 2
        return 0

    out = Path(args.folder) / "analysis" / REPORT_SUBDIR
    out.mkdir(parents=True, exist_ok=True)
    (out / "report_data.json").write_text(json.dumps(report, separators=(",", ":")), encoding="utf-8")
    write_html(report, out / "report.html")
    if not args.no_xlsx:
        write_workbook(report, out / "calibration_report.xlsx")
    if not args.no_png:
        write_heatmap_pngs(report, out)
    print(f"Report written to {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
